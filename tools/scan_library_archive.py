"""Read saved List Views and Git history; produce a cover-free recovery report.

This does not change the app archive, original workbooks, or assign any Book IDs.
"""
from __future__ import annotations
import argparse
from collections import defaultdict
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import unicodedata

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SHARED = Path(r'C:\Users\cjade\OneDrive\Shared Workbooks\MyLibrary')
CURRENT = SHARED / 'LIBRARY LIST VIEW.xlsx'
OLD_ARCHIVE = ROOT.parent / 'MyLibrary Archive' / 'ARCHIVE.xlsx'
BASE_HEADERS = ['CJ', 'JC', 'BIPOC', 'LGBTQ+', 'ISBN', 'Year', 'Pages', 'Title',
                'Series', 'First', 'Last', 'Genre', 'Subgenre', 'Publisher', 'Origin',
                'Bookcase', 'Shelf', 'Position', 'SYSTEM COLUMNS - AUTOMATION ONLY',
                'Needs Review', 'Book ID', 'Series Sort', 'Volume Sort', 'Last Sort', 'First Sort']
JSON_FIELDS = {'cj': 'CJ', 'jc': 'JC', 'bipoc': 'BIPOC', 'lgbtq': 'LGBTQ+',
               'isbn': 'ISBN', 'publicationYear': 'Year', 'totalPages': 'Pages',
               'rawTitle': 'Title', 'series': 'Series', 'firstName': 'First',
               'lastName': 'Last', 'genre': 'Genre', 'subgenre': 'Subgenre',
               'publisher': 'Publisher', 'origin': 'Origin', 'bookcase': 'Bookcase',
               'rawShelf': 'Shelf', 'shelfPosition': 'Position', 'bookId': 'Book ID',
               'notes': 'Notes', 'archiveStatus': 'Prior archive status',
               'archiveNotes': 'Archive Notes', 'offloadDate': 'Offload Date',
               'offloadDestination': 'Offload Destination'}


def text(value):
    return '' if value is None else str(value).strip()


def normalized(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', text(value)).casefold() if c.isalnum())


def stable_id(row):
    value = text(row.get('Book ID')).lower()
    return value if re.fullmatch(r'book-[0-9a-f]{8}(?:-[0-9a-f]{4}){3}-[0-9a-f]{12}', value) else ''


def isbn(row):
    return re.sub(r'[^0-9X]', '', text(row.get('ISBN')).upper())


def title_author(row):
    # Older exports included series annotations in raw titles, e.g. (Series #3).
    title = re.sub(r'\s*\([^()]*#\s*\d[^()]*\)\s*$', '', text(row.get('Title')))
    return (normalized(title), normalized(text(row.get('First')) + ' ' + text(row.get('Last'))))


def fingerprint(row):
    return (*title_author(row), isbn(row))


def stamp(path):
    match = re.search(r'(20\d{6})-(\d{6})', path.name)
    if match:
        return datetime.strptime(''.join(match.groups()), '%Y%m%d%H%M%S').isoformat()
    return datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec='seconds')


def clean_cell(value):
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def read_xlsx(path, existing_archive=False):
    # Read a disposable snapshot; never save the source through openpyxl.
    with tempfile.TemporaryDirectory() as folder:
        copy = Path(folder) / path.name
        shutil.copy2(path, copy)
        digest = hashlib.sha256(copy.read_bytes()).hexdigest()
        book = load_workbook(copy, read_only=True, data_only=True)
        sheet_name = 'Archive' if existing_archive else 'List View'
        if sheet_name not in book.sheetnames:
            book.close()
            return [], [], digest
        sheet = book[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = None
        header_row = 0
        for number, values in enumerate(rows, 1):
            if 'Title' in values and ('ISBN' in values or 'Book ID' in values):
                headers = [text(v) for v in values]
                header_row = number
                break
            if number >= 30:
                break
        if headers is None:
            book.close()
            raise ValueError(f'No List View headers found: {path}')
        renames = {'Publication Year': 'Year', 'Total Pages': 'Pages', 'CJ Read': 'CJ',
                   'JC Read': 'JC', 'Last Known Bookcase': 'Bookcase', 'Raw Shelf': 'Shelf',
                   'Original Notes': 'Notes', 'Archive Status': 'Prior archive status'}
        excluded = {'Local Cover File', 'App Cover Path', 'Format', 'Catalog Key'}
        selected = [(i, renames.get(h, h) if existing_archive else h)
                    for i, h in enumerate(headers) if h and h not in excluded]
        records = []
        for number, values in enumerate(rows, header_row + 1):
            row = {h: clean_cell(values[i]) if i < len(values) else None for i, h in selected}
            if text(row.get('Title')):
                if row.get('ISBN') is not None:
                    row['ISBN'] = text(row['ISBN'])
                records.append({'values': row, 'source': str(path), 'sheet': sheet_name,
                                'row': number, 'date': stamp(path), 'raw': not existing_archive})
        book.close()
        return records, [h for _, h in selected], digest


def json_observations(data, source, date):
    records = []
    if not isinstance(data, list):
        raise ValueError(f'Expected a book array in {source}')
    for number, book in enumerate(data, 1):
        row = {header: clean_cell(book[key]) for key, header in JSON_FIELDS.items() if key in book}
        row.setdefault('Title', book.get('title'))
        if not text(row.get('Title')):
            continue
        row['ISBN'] = text(row.get('ISBN'))
        records.append({'values': row, 'source': source, 'sheet': 'book records',
                        'row': number, 'date': date, 'raw': False})
    return records


def classify(observations, current):
    """IDs are authoritative; legacy links require a unique exact fingerprint."""
    current_ids = {stable_id(r) for r in current if stable_id(r)}
    current_isbn = defaultdict(list)
    current_title = defaultdict(list)
    for index, row in enumerate(current):
        if isbn(row):
            current_isbn[isbn(row)].append(index)
        current_title[title_author(row)].append(index)
    known = defaultdict(set)
    known_title = defaultdict(set)
    for row in [*(o['values'] for o in observations), *current]:
        if stable_id(row):
            known[fingerprint(row)].add(stable_id(row))
            known_title[title_author(row)].add(stable_id(row))
    def inferred_id(row):
        candidates = known[fingerprint(row)]
        if not isbn(row):
            candidates = known_title[title_author(row)]
        return next(iter(candidates)) if len(candidates) == 1 else ''
    legacy_links = defaultdict(set)
    for obs in observations:
        row = obs['values']
        if not stable_id(row) and text(row.get('Book ID')) and inferred_id(row):
            legacy_links[text(row['Book ID'])].add(inferred_id(row))
    groups = defaultdict(list)
    for obs in observations:
        row = obs['values']
        key = stable_id(row)
        if not key:
            key = inferred_id(row)
        if not key and len(legacy_links[text(row.get('Book ID'))]) == 1:
            key = next(iter(legacy_links[text(row.get('Book ID'))]))
        if not key:
            key = 'legacy-' + hashlib.sha256(json.dumps(fingerprint(row)).encode()).hexdigest()[:16]
        groups[key].append(obs)
    result = []
    for key, versions in groups.items():
        if key in current_ids:
            continue
        values = {}
        ordered = sorted(versions, key=lambda o: (o['raw'], o['date'], o['source'], o['row']))
        for obs in ordered:
            values.update(obs['values'])
        # Prefer the stable ID even if the latest source is a legacy blank-ID row.
        if key.startswith('book-'):
            values['Book ID'] = key
        possible_indices = set()
        for version in versions:
            possible_indices.update(current_isbn[isbn(version['values'])])
            for index in current_title[title_author(version['values'])]:
                # A known different edition is not evidence that this copy remains.
                if not isbn(values) or not isbn(current[index]) or isbn(values) == isbn(current[index]):
                    possible_indices.add(index)
        possible = [current[i] for i in sorted(possible_indices)]
        status = 'Possible current match' if possible else ('Missing from current list' if key.startswith('book-') else 'Legacy identity needs review')
        result.append({'key': key, 'values': values, 'status': status,
                       'firstSeen': min(v['date'] for v in versions),
                       'lastSeen': max(v['date'] for v in versions),
                       'preferred': ordered[-1], 'versions': versions,
                       'possible': '; '.join(f"{r['Title']} [{text(r.get('Book ID'))}]" for r in possible)})
    return sorted(result, key=lambda r: (normalized(r['values'].get('Last')), normalized(r['values'].get('Title')), r['key']))


def scan(current_path, history_dirs, include_git=True):
    current_obs, headers, current_hash = read_xlsx(current_path)
    if not current_obs:
        raise ValueError('Current List View is empty; refusing to mark everything missing.')
    sources = [{'source': str(current_path), 'kind': 'Current List View', 'date': stamp(current_path),
                'rows': len(current_obs), 'sha256': current_hash}]
    observations = []
    paths = set()
    for folder in history_dirs:
        for path in folder.rglob('*.xlsx'):
            if not path.name.startswith('~$') and ('list view' in path.name.lower() or path.name.startswith('LIBRARY-before-book-id')):
                if path.resolve() != current_path.resolve():
                    paths.add(path.resolve())
    if OLD_ARCHIVE.exists():
        paths.add(OLD_ARCHIVE.resolve())
    for path in sorted(paths):
        rows, found_headers, digest = read_xlsx(path, path == OLD_ARCHIVE.resolve())
        if not rows:
            continue
        observations.extend(rows)
        headers += [h for h in found_headers if h not in headers]
        sources.append({'source': str(path), 'kind': 'Saved workbook', 'date': stamp(path),
                        'rows': len(rows), 'sha256': digest})
    for archive in [ROOT / 'public/data/library-archive.json', *sorted((ROOT / 'tools/recovery').glob('*.json'))]:
        if not archive.exists():
            continue
        data = json.loads(archive.read_text(encoding='utf-8-sig'))
        rows = json_observations(data, str(archive), stamp(archive))
        observations.extend(rows)
        sources.append({'source': str(archive), 'kind': 'App archive', 'date': stamp(archive),
                        'rows': len(rows), 'sha256': hashlib.sha256(archive.read_bytes()).hexdigest()})
    if include_git:
        log = subprocess.check_output(['git', 'log', '--format=%H %aI', '--', 'public/data/library-books.json'], cwd=ROOT, text=True)
        for line in log.splitlines():
            commit, date = line.split(' ', 1)
            raw = subprocess.check_output(['git', 'show', f'{commit}:public/data/library-books.json'], cwd=ROOT)
            # Early commits can precede permanent IDs, so preserve them as legacy evidence.
            rows = json_observations(json.loads(raw), f'git:{commit}:public/data/library-books.json', date[:19])
            observations.extend(rows)
            sources.append({'source': f'git:{commit}:public/data/library-books.json', 'kind': 'Git book snapshot',
                            'date': date[:19], 'rows': len(rows), 'sha256': hashlib.sha256(raw).hexdigest()})
    for obs in observations:
        headers += [h for h in obs['values'] if h not in headers]
    records = classify(observations, [o['values'] for o in current_obs])
    return {'generatedAt': datetime.now(timezone.utc).isoformat(), 'headers': headers,
            'currentCount': len(current_obs), 'sources': sources, 'records': records,
            'historyRowsRead': len(observations)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--current', type=Path, default=CURRENT)
    parser.add_argument('--history-dir', type=Path, action='append', default=[])
    parser.add_argument('--no-git', action='store_true')
    parser.add_argument('--output', type=Path, required=True, help='Intermediate recovery JSON (not the app archive)')
    args = parser.parse_args()
    payload = scan(args.current, [ROOT / 'tools', *args.history_dir], not args.no_git)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    from collections import Counter
    print(json.dumps({'sources': len(payload['sources']), 'historicalRows': payload['historyRowsRead'],
                      'currentBooks': payload['currentCount'], 'results': dict(Counter(r['status'] for r in payload['records']))}, indent=2))


if __name__ == '__main__':
    main()
