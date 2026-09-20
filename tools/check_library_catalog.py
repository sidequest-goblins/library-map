"""Read saved workbooks and print a shareable catalog-completion checklist."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
import re
import sys
import warnings

from openpyxl import load_workbook
from build_library_data import (
    CATALOG_WORKBOOK_PATH, LIST_VIEW_WORKBOOK_PATH, clean, load_catalog_books,
    make_author, make_author_sort, make_catalog_key, normalize_header,
    normalize_match_text, is_short_title_match, parse_title,
)


def read_list(path):
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if "List View" not in workbook.sheetnames:
            raise ValueError("List View worksheet is missing.")
        rows = workbook["List View"].iter_rows(values_only=True)
        headers = [normalize_header(v) for v in next(rows, ())]
        if not {"title", "first", "last"}.issubset(headers):
            raise ValueError("List View needs Title, First, and Last headers in row 1.")
        books = []
        for number, row in enumerate(rows, 2):
            values = dict(zip(headers, row))
            raw = clean(values.get("title"))
            if not raw:
                continue
            first, last = clean(values.get("first")), clean(values.get("last"))
            books.append(dict(
                title=clean(parse_title(raw)["title"]), rawTitle=raw,
                author=make_author(first, last), authorSort=make_author_sort(first, last),
                bookId=clean(values.get("bookid")), row=number,
                location=", ".join(clean(values.get(k)) for k in ("bookcase", "shelf") if clean(values.get(k))),
            ))
        if not books:
            raise ValueError("No books found on List View; comparison stopped.")
        return books
    finally:
        workbook.close()


def classify(book, catalog, duplicates):
    """Exact app keys count as present; possible/ambiguous matches remain visible."""
    key = make_catalog_key(book["title"], book["authorSort"])
    if not book["author"]:
        return "review", "List View author is blank", []
    if key in duplicates:
        return "review", "Multiple catalog rows share this title/author key", duplicates[key]
    if key in catalog:
        candidate = catalog[key]
        # The app key uses only the first contributor. Do not hide contributor differences.
        if normalize_match_text(book["author"]) != normalize_match_text(candidate["author"]):
            return "review", "Title matches but contributor names differ", [candidate]
        return "matched", "Exact title/author", [candidate]
    candidates = [b for b in catalog.values()
                  if normalize_match_text(b["authorSort"]) == normalize_match_text(book["authorSort"])
                  and (is_short_title_match(book["title"], b["title"])
                       or is_short_title_match(b["title"], book["title"]))]
    if candidates:
        return "review", "Possible shortened title or different volume/edition", candidates
    candidates = [b for b in catalog.values()
                  if normalize_match_text(b["title"]) == normalize_match_text(book["title"])]
    if candidates:
        return "review", "Same title with different author spelling or attribution", candidates
    # Suggestions are never accepted automatically. Keep different volume numbers apart.
    title = normalize_match_text(book["title"])
    ranked = sorted(
        ((SequenceMatcher(None, title, normalize_match_text(b["title"])).ratio(), b)
         for b in catalog.values()
         if re.findall(r"\d+", title) == re.findall(r"\d+", normalize_match_text(b["title"]))),
        key=lambda pair: pair[0], reverse=True,
    )
    candidates = [b for score, b in ranked[:3] if score >= 0.88]
    if candidates:
        return "review", "Similar title; verify title and author manually", candidates
    return "missing", "No catalog match found", []


def compare(books, catalog, duplicates):
    results = []
    for book in sorted(books, key=lambda b: (b["authorSort"].casefold(), b["rawTitle"].casefold(), b["row"])):
        status, reason, candidates = classify(book, catalog, duplicates)
        results.append(dict(book=book, status=status, reason=reason, candidates=candidates))
    return results


def render_report(results, catalog_report, list_path, catalog_path):
    counts = Counter(r["status"] for r in results)
    lines = ["LIBRARY CATALOG CHECK", "=" * 60,
             f"Checked: {datetime.now().astimezone():%Y-%m-%d %I:%M %p %Z}",
             f"List: {list_path.name} / List View", f"Catalog: {catalog_path.name}",
             f"List View books: {len(results)} | Catalog rows: {catalog_report['catalogRowsWithTitles']}",
             f"Matched: {counts['matched']} | No match found: {counts['missing']} | Review: {counts['review']}", "",
             "Saved workbook contents only. Save Excel changes before rerunning.",
             "Title/author comparison; LIBRARY.xlsx has no shared Book ID or ISBN column.",
             "No match found means likely missing; check spelling before adding a duplicate.",
             "To Buy, To Trade, and Series to Complete are excluded.",
             "Catalog sheets checked: " + ", ".join(catalog_report["processedCatalogSheets"]),
             "Catalog sheets skipped (no catalog headers): " + ", ".join(catalog_report["skippedCatalogSheets"]), ""]
    for status, heading in (("missing", "BOOKS WITH NO CATALOG MATCH"), ("review", "CHECK POSSIBLE MATCHES BEFORE ADDING")):
        entries = [r for r in results if r["status"] == status]
        lines += [f"{heading} ({len(entries)})", "-" * 60]
        if not entries:
            lines.append("None.")
        for item in entries:
            book = item["book"]
            # Collapse embedded newlines for a readable one-book-per-line checklist.
            title, author = " ".join(book["rawTitle"].split()), " ".join(book["author"].split())
            lines.append(f"[ ] {title} — {author or '(author missing)'}")
            lines.append(f"    List View row {book['row']}" + (f" | {book['location']}" if book["location"] else ""))
            if status == "review":
                lines.append(f"    {item['reason']}")
                for candidate in item["candidates"]:
                    lines.append(f"    Catalog: {' '.join(candidate['rawTitle'].split())} — {' '.join(candidate['author'].split())} ({candidate['sourceSheet']}, row {candidate['sourceRow']})")
            lines.append("")
        lines.append("")
    lines += ["Rerun after catalog updates to refresh this checklist."]
    return "\n".join(lines) + "\n"


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--list-view", type=Path, default=LIST_VIEW_WORKBOOK_PATH)
    parser.add_argument("--catalog", type=Path, default=CATALOG_WORKBOOK_PATH)
    parser.add_argument("--output", type=Path, help="UTF-8 text report; defaults beside List View")
    args = parser.parse_args(argv)
    output = args.output or args.list_view.parent / "MISSING FROM CATALOG.txt"
    if output.suffix.lower() != ".txt":
        parser.error("Output must be a .txt file.")
    if output.resolve() in (args.list_view.resolve(), args.catalog.resolve()):
        parser.error("Output must not overwrite an input workbook.")
    try:
        # This warning refers to saving through openpyxl. This tool never saves a workbook.
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported and will be removed", category=UserWarning)
        before = [(p.stat().st_size, p.stat().st_mtime_ns) for p in (args.list_view, args.catalog)]
        books = read_list(args.list_view)
        workbook = load_workbook(args.catalog, read_only=True, data_only=True, keep_links=False)
        try:
            catalog, metadata = load_catalog_books(workbook)
        finally:
            workbook.close()
        if not catalog:
            raise ValueError("No catalog books found. Expected Cover, Title, First, Last, Format headers.")
        after = [(p.stat().st_size, p.stat().st_mtime_ns) for p in (args.list_view, args.catalog)]
        if before != after:
            raise ValueError("A workbook changed during comparison. Wait for syncing to finish and rerun.")
        report = render_report(compare(books, catalog, metadata["duplicateCatalogKeys"]), metadata, args.list_view, args.catalog)
        # Atomic replacement keeps the previous report intact if writing fails.
        import tempfile
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8-sig", dir=output.parent, delete=False, suffix=".tmp") as handle:
                temporary = Path(handle.name)
                handle.write(report)
            temporary.replace(output)
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
        print(report)
        print(f"Saved checklist: {output}")
        return 0
    except (OSError, ValueError) as exc:
        print(f"Catalog check failed: {exc}\nNo new report was saved; any previous checklist is out of date.", file=sys.stderr)
        return 1


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    raise SystemExit(main())
